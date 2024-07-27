import os
import subprocess
import extract_msg
from pdfminer.pdfparser import PDFParser
from pdfminer.pdfdocument import PDFDocument
from pdfminer.pdftypes import resolve1
from PIL import Image
from PIL.ExifTags import TAGS
from docx import Document
import openpyxl
import pandas as pd
from datetime import datetime

def format_pdf_date(pdf_date):
    if pdf_date and isinstance(pdf_date, bytes):
        pdf_date = pdf_date.decode('utf-8')
        if pdf_date.startswith('D:'):
            pdf_date = pdf_date[2:]
            return f"{pdf_date[4:6]}/{pdf_date[6:8]}/{pdf_date[:4]} {pdf_date[8:10]}:{pdf_date[10:12]}:{pdf_date[12:14]}"
    return None

def get_pdf_dates(pdf_path):
    try:
        with open(pdf_path, 'rb') as file:
            parser = PDFParser(file)
            document = PDFDocument(parser)
            if not document.info:
                return None, None
            metadata = resolve1(document.info[0])
            creation_date = resolve1(metadata.get('CreationDate')) if 'CreationDate' in metadata else None
            modification_date = resolve1(metadata.get('ModDate')) if 'ModDate' in metadata else None
    except Exception as e:
        print(f"Error reading {pdf_path}: {e}")
        return None, None

    formatted_creation_date = format_pdf_date(creation_date)
    formatted_modification_date = format_pdf_date(modification_date)

    return formatted_creation_date, formatted_modification_date

def format_image_date(date_str):
    try:
        if date_str:
            dt = datetime.strptime(date_str, '%Y:%m:%d %H:%M:%S')
            return dt.strftime('%m/%d/%Y %H:%M:%S')
    except Exception as e:
        print(f"Error formatting image date: {e}")
        return None

def get_image_dates(image_path):
    try:
        image = Image.open(image_path)
        exif_data = image._getexif()
        if not exif_data:
            return None, None

        exif = {TAGS.get(tag, tag): value for tag, value in exif_data.items()}
        creation_date = exif.get('DateTimeOriginal')
        modification_date = exif.get('DateTime')
        formatted_creation_date = format_image_date(creation_date)
        formatted_modification_date = format_image_date(modification_date)
    except Exception as e:
        print(f"Error reading {image_path}: {e}")
        return None, None

    return formatted_creation_date, formatted_modification_date

def format_doc_date(date):
    try:
        if date:
            dt = datetime.fromisoformat(str(date))
            return dt.strftime('%m/%d/%Y %H:%M:%S')
    except Exception as e:
        print(f"Error formatting document date: {e}")
        return None

def convert_doc_to_docx(doc_path):
    try:
        subprocess.call(['libreoffice', '--headless', '--convert-to', 'docx', doc_path, '--outdir', os.path.dirname(doc_path)])
        return doc_path.replace('.doc', '.docx')
    except Exception as e:
        print(f"Error converting {doc_path} to DOCX: {e}")
        return None

def get_docx_dates(docx_path):
    try:
        doc = Document(docx_path)
        core_properties = doc.core_properties
        creation_date = format_doc_date(core_properties.created)
        modification_date = format_doc_date(core_properties.modified)
    except Exception as e:
        print(f"Error reading {docx_path}: {e}")
        return None, None

    return creation_date, modification_date

def get_excel_dates(excel_path):
    try:
        workbook = openpyxl.load_workbook(excel_path, read_only=True)
        props = workbook.properties
        creation_date = format_doc_date(props.created)
        modification_date = format_doc_date(props.modified)
    except Exception as e:
        print(f"Error reading {excel_path}: {e}")
        return None, None

    return creation_date, modification_date

def get_msg_info(msg_path):
    try:
        msg = extract_msg.Message(msg_path)
        from_email = msg.sender
        to_email = msg.to
        cc_email = msg.cc
    except Exception as e:
        print(f"Error reading {msg_path}: {e}")
        return None, None, None

    return from_email, to_email, cc_email

def read_files_from_folder(folder_path):
    file_data = []
    base_path_length = len(folder_path) + 1  # To remove the base folder path from the full path

    for root, dirs, files in os.walk(folder_path):
        for file in files:
            file_path = os.path.join(root, file)
            relative_path = os.path.dirname(file_path[base_path_length:])
            folder_name = os.path.basename(relative_path)
            creation_date, modification_date = None, None
            from_email, to_email, cc_email = None, None, None
            file_type = 'Other'

            if file.lower().endswith('.pdf'):
                creation_date, modification_date = get_pdf_dates(file_path)
                file_type = 'PDF'
            elif file.lower().endswith('.jpg') or file.lower().endswith('.jpeg'):
                creation_date, modification_date = get_image_dates(file_path)
                file_type = 'Image'
            elif file.lower().endswith('.docx'):
                creation_date, modification_date = get_docx_dates(file_path)
                file_type = 'Word'
            elif file.lower().endswith('.doc'):
                docx_path = convert_doc_to_docx(file_path)
                if docx_path:
                    creation_date, modification_date = get_docx_dates(docx_path)
                    file_type = 'Word'
            elif file.lower().endswith('.xlsx'):
                creation_date, modification_date = get_excel_dates(file_path)
                file_type = 'Excel'
            elif file.lower().endswith('.msg'):
                from_email, to_email, cc_email = get_msg_info(file_path)
                file_type = 'Email'

            file_data.append({
                'Filename': file,
                'Creation Date': creation_date,
                'Modification Date': modification_date,
                'File Type': file_type,
                'Path': relative_path,
                'Folder Name': folder_name,
                'From Email': from_email,
                'To Email': to_email,
                'CC Email': cc_email
            })

    return file_data

folder_path = 'path_to_your_folder'
file_data = read_files_from_folder(folder_path)

# Create a DataFrame and save to Excel
df = pd.DataFrame(file_data)
output_file = 'file_metadata.xlsx'
df.to_excel(output_file, index=False)

print(f"Metadata extracted and saved to {output_file}")
